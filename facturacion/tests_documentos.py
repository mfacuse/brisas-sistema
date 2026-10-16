"""Numeración interna, notas de cobro en PDF, XML sin notas de cobro y vínculo con las facturas emitidas por BSOFT."""
import io
import shutil
import tempfile
from datetime import date
from decimal import Decimal as D
from pathlib import Path

from django.contrib.auth.models import Group, User
from django.test import TestCase, override_settings
from fpdf import FPDF
from pypdf import PdfReader

from indicadores.models import UF
from maestro.models import Arrendatario, Contrato, Propiedad, Sociedad, TramoRenta

from . import bsoft, emitidas, nota_cobro, servicios
from .models import CodigoServicioBsoft, Documento, FacturaEmitida

TEXTO_FACTURA = """INMOBILIARIA BOULEVARD TALCA SPA
Giro: Arriendo Inmuebles Amoblados
R.U.T.: 76.000.000-0
FACTURA ELECTRONICA
Nº {folio}
S.I.I. - TALCA
Fecha Emision: 05 de Octubre del 2026
SEÑOR(ES): {cliente}
R.U.T.: {rut}
Cantidad Detalle Valor Total
1 ARRIENDO 1.108.765
MONTO NETO $ {neto}
I.V.A. 19% $ {iva}
TOTAL $ {total}
"""


def _pdf_facturas(ruta: Path, facturas):
    """PDF de prueba con una factura por página (texto parecido al de un DTE del SII)."""
    pdf = FPDF()
    pdf.set_font('Helvetica', size=10)
    for f in facturas:
        pdf.add_page()
        for linea in TEXTO_FACTURA.format(**f).splitlines():
            pdf.cell(0, 6, linea.encode('latin-1', 'replace').decode('latin-1'), new_x='LMARGIN', new_y='NEXT')
    ruta.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(ruta))


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.soc = Sociedad.objects.create(nombre='Inmobiliaria Prueba SpA', rut='76.000.000-0', sigla='IP',
                                          datos_pago='Banco Prueba · Cta. Cte. 123-456 · pagos@prueba.cl')
        prop = Propiedad.objects.create(sociedad=cls.soc, codigo='prueba', nombre='Centro Prueba', nombre_glosa='Centro Prueba Talca')
        cls.arr = Arrendatario.objects.create(razon_social='Ale SpA', rut='11.111.111-1')
        cls.arr2 = Arrendatario.objects.create(razon_social='Casa Ltda', rut='22.222.222-2')
        cls.fa = Contrato.objects.create(codigo='P-01', propiedad=prop, sociedad=cls.soc, arrendatario=cls.arr, nombre_fantasia='Ale',
                                         inicio_vigencia=date(2025, 1, 1))
        cls.nc = Contrato.objects.create(codigo='P-02', propiedad=prop, sociedad=cls.soc, arrendatario=cls.arr2, nombre_fantasia='Casa',
                                         inicio_vigencia=date(2025, 1, 1), tipo_documento='NC', direccion_contrato='1 Norte 1209, Talca')
        for c, m in ((cls.fa, '27'), (cls.nc, '10')):
            TramoRenta.objects.create(contrato=c, desde=date(2025, 1, 1), monto=D(m), unidad='UF')
        UF.objects.create(fecha=date(2026, 10, 1), valor=D('41065.38'))
        UF.objects.create(fecha=date(2026, 11, 1), valor=D('41252.94'))
        for o in ('RENTA', 'RDIF', 'DESC', 'ELEC', 'GAS', 'OTRO'):
            CodigoServicioBsoft.objects.create(sociedad=cls.soc, origen=o, codigo='1002')

    def generar_y_aprobar(self, mes=10):
        servicios.generar_borradores(self.soc, 2026, mes)
        servicios.aprobar(self.soc, 2026, mes, 'op')
        return {d.contrato.codigo: d for d in Documento.objects.filter(periodo=date(2026, mes, 1))}


class NumeracionTest(Base):
    def test_numero_correlativo_por_sociedad_y_serie(self):
        d10 = self.generar_y_aprobar(10)
        self.assertEqual(d10['P-01'].numero_interno, 'IP-F-0001')
        self.assertEqual(d10['P-02'].numero_interno, 'IP-NC-0001')
        self.assertEqual((d10['P-02'].iva, d10['P-02'].total), (D('0'), d10['P-02'].exento))  # nota de cobro: sin IVA
        d11 = self.generar_y_aprobar(11)
        self.assertEqual((d11['P-01'].numero_interno, d11['P-02'].numero_interno), ('IP-F-0002', 'IP-NC-0002'))

    def test_volver_a_borrador_y_regenerar_conserva_el_numero(self):
        d = self.generar_y_aprobar(10)['P-01']
        Documento.objects.filter(pk=d.pk).update(estado='B')
        servicios.generar_borradores(self.soc, 2026, 10)
        nuevo = Documento.objects.get(contrato=self.fa, periodo=date(2026, 10, 1))
        self.assertEqual(nuevo.numero_interno, 'IP-F-0001')
        servicios.aprobar(self.soc, 2026, 10, 'op')
        nuevo.refresh_from_db()
        self.assertEqual(nuevo.numero_interno, 'IP-F-0001')

    def test_sigla_por_defecto(self):
        self.assertEqual(servicios.sigla(Sociedad(nombre='Inmobiliaria Boulevard Talca SpA', pk=9)), 'BT')
        self.assertEqual(servicios.sigla(Sociedad(nombre='SpA', pk=9)), 'S9')


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='brisas_salidas_'))
class NotaCobroTest(Base):
    def test_xml_de_bsoft_no_incluye_notas_de_cobro(self):
        self.generar_y_aprobar(10)
        _, _, archivos = bsoft.exportar_xml(self.soc, 2026, 10)
        self.assertEqual([Path(a).name for a in archivos], ['P-01_2026-10.xml'])
        Documento.objects.filter(contrato=self.fa).delete()
        with self.assertRaisesMessage(ValueError, 'las notas de cobro no van a BSOFT'):
            bsoft.exportar_xml(self.soc, 2026, 10)

    def test_pdf_de_la_nota_de_cobro(self):
        d = self.generar_y_aprobar(10)['P-02']
        datos = nota_cobro.pdf_nota(d)
        self.assertTrue(datos.startswith(b'%PDF'))
        texto = PdfReader(io.BytesIO(datos)).pages[0].extract_text()
        for esperado in ('NOTA DE COBRO', 'IP-NC-0001', 'Casa Ltda', '22.222.222-2', 'Octubre 2026', 'Banco Prueba', 'no tributario'):
            self.assertIn(esperado, texto)
        with self.assertRaises(ValueError):
            nota_cobro.pdf_nota(self.generar_y_aprobar(10)['P-01'])  # una factura no es nota de cobro

    def test_borrador_sale_marcado(self):
        servicios.generar_borradores(self.soc, 2026, 10)
        d = Documento.objects.get(contrato=self.nc)
        self.assertIn('BORRADOR', PdfReader(io.BytesIO(nota_cobro.pdf_nota(d))).pages[0].extract_text())

    def test_panel_descarga_notas_del_mes_y_las_marca_emitidas(self):
        self.generar_y_aprobar(10)
        Group.objects.create(name='Operación')
        u = User.objects.create_user('op', password='x')
        u.groups.add(Group.objects.get(name='Operación'))
        self.client.force_login(u)
        r = self.client.get(f'/facturacion/notas-cobro/?sociedad={self.soc.id}&periodo=2026-10')
        self.assertEqual(r['Content-Type'], 'application/zip')
        self.assertEqual(Documento.objects.get(contrato=self.nc).estado, 'X')
        self.assertEqual(Documento.objects.get(contrato=self.fa).estado, 'A')  # la factura espera el XML
        d = Documento.objects.get(contrato=self.nc)
        self.assertEqual(self.client.get(f'/documento/{d.pk}/nota-cobro.pdf')['Content-Type'], 'application/pdf')
        self.assertContains(self.client.get(f'/facturacion/?sociedad={self.soc.id}&periodo=2026-10'), 'IP-NC-0001')


class LecturaFacturaTest(TestCase):
    def test_lee_folio_rut_fecha_y_montos(self):
        t = TEXTO_FACTURA.format(folio='5148', cliente='Ale SpA', rut='11.111.111-1', neto='1.108.765', iva='210.665', total='1.319.430')
        d = emitidas.leer_texto(t, '76.000.000-0')
        self.assertEqual((d.tipo_dte, d.folio, d.rut_receptor, d.fecha_emision), ('33', '5148', '11111111-1', date(2026, 10, 5)))
        self.assertEqual((d.neto, d.iva, d.total), (D('1108765'), D('210665'), D('1319430')))

    def test_exenta_y_fecha_numerica(self):
        t = 'FACTURA NO AFECTA O EXENTA ELECTRONICA\nN° 000321\nR.U.T.: 76.000.000-0\nFecha: 01/11/2026\nRUT: 22.222.222-2\nMONTO EXENTO $ 410.529\nTOTAL $ 410.529'
        d = emitidas.leer_texto(t, '76.000.000-0')
        self.assertEqual((d.tipo_dte, d.folio, d.rut_receptor, d.fecha_emision, d.total), ('34', '321', '22222222-2', date(2026, 11, 1), D('410529')))


class VinculoTest(Base):
    def setUp(self):
        self.raiz = Path(tempfile.mkdtemp(prefix='brisas_emitidas_')).resolve()
        self.addCleanup(shutil.rmtree, self.raiz, ignore_errors=True)
        o = override_settings(CARPETAS_RAIZ=[self.raiz])
        o.enable()
        self.addCleanup(o.disable)
        self.soc.carpeta_facturas = str(self.raiz)
        self.soc.save()
        self.docs = self.generar_y_aprobar(10)

    def _f(self, folio, rut, total, cliente='Cliente'):
        return dict(folio=folio, cliente=cliente, rut=rut, neto='1', iva='1', total=f'{int(total):,}'.replace(',', '.'))

    def test_vincula_por_rut_y_total_y_deja_lo_dudoso_para_revisar(self):
        fa = self.docs['P-01']
        # la 5149 llega primero: mismo RUT pero otro monto → para revisar; la 5148 calza exacto → vinculada
        _pdf_facturas(self.raiz / 'Octubre' / 'facturas.pdf', [self._f('5149', '11.111.111-1', 999), self._f('5148', '11.111.111-1', fa.total),
                                                               self._f('5150', '99.999.999-9', 1000)])
        r = emitidas.revisar(self.soc)
        self.assertEqual((r.facturas, r.vinculadas, r.dudosas, r.sin_documento), (3, 1, 1, 1))
        fa.refresh_from_db()
        self.assertEqual((fa.folio, fa.numero_interno, fa.emitida.folio), ('5148', 'IP-F-0001', '5148'))
        dudosa = FacturaEmitida.objects.get(folio='5149')
        self.assertEqual(dudosa.estado, 'D')
        # una segunda lectura no repite nada
        self.assertEqual(emitidas.revisar(self.soc).archivos_leidos, 0)
        # a mano: se desvincula la buena y se vincula la dudosa
        emitidas.desvincular(fa.emitida, 'op')
        emitidas.vincular_manual(dudosa, fa, 'op')
        fa.refresh_from_db()
        self.assertEqual((fa.folio, fa.emitida.estado), ('5149', 'M'))
        with self.assertRaises(ValueError):
            emitidas.vincular_manual(FacturaEmitida.objects.get(folio='5150'), fa, 'op')

    def test_la_nota_de_cobro_no_se_vincula(self):
        nc = self.docs['P-02']
        _pdf_facturas(self.raiz / 'x.pdf', [self._f('77', '22.222.222-2', nc.total)])
        emitidas.revisar(self.soc)
        self.assertEqual(FacturaEmitida.objects.get().estado, 'S')

    def test_pdf_ilegible(self):
        (self.raiz / 'escaneo.pdf').write_bytes(b'no es un pdf')
        r = emitidas.revisar(self.soc)
        self.assertEqual(r.ilegibles, 1)
        self.assertEqual(FacturaEmitida.objects.get().estado, 'I')

    def test_pantalla_de_facturas_emitidas(self):
        _pdf_facturas(self.raiz / 'f.pdf', [self._f('5148', '11.111.111-1', self.docs['P-01'].total)])
        Group.objects.create(name='Operación')
        u = User.objects.create_user('op', password='x')
        u.groups.add(Group.objects.get(name='Operación'))
        self.client.force_login(u)
        self.client.post('/facturacion/emitidas/revisar/', {'sociedad': self.soc.id})
        r = self.client.get(f'/facturacion/emitidas/?sociedad={self.soc.id}')
        self.assertContains(r, '5148')
        self.assertContains(r, 'IP-F-0001')
        f = FacturaEmitida.objects.get()
        self.assertEqual(self.client.get(f'/facturacion/emitidas/{f.pk}/pdf/')['Content-Type'], 'application/pdf')
        self.assertContains(self.client.get(f'/documento/{self.docs["P-01"].pk}/'), 'Folio SII')


class MedidoresSinContratosTest(TestCase):
    def test_sembrar_medidores_solo_para_contratos_existentes(self):
        from . import cargas
        from .models import MedidorServicio
        self.assertEqual(cargas.sembrar_medidores(), 0)  # sistema vacío: no se cae
        soc = Sociedad.objects.create(nombre='Inmobiliaria Prueba SpA', rut='76.000.000-0')
        prop = Propiedad.objects.create(sociedad=soc, codigo='prueba', nombre='Centro Prueba')
        arr = Arrendatario.objects.create(razon_social='Ale SpA', rut='11.111.111-1')
        Contrato.objects.create(codigo='B-02', propiedad=prop, sociedad=soc, arrendatario=arr, nombre_fantasia='Ale', inicio_vigencia=date(2025, 1, 1))
        n = cargas.sembrar_medidores()
        self.assertGreater(n, 0)
        self.assertEqual(set(MedidorServicio.objects.values_list('contrato__codigo', flat=True)), {'B-02'})
        self.assertEqual(cargas.sembrar_medidores(), 0)  # no duplica


class RenovacionAutomaticaTest(TestCase):
    def test_sigue_la_renta_del_ultimo_tramo(self):
        from . import motor
        t = [motor.TramoD(date(2024, 1, 1), date(2025, 12, 31), D('20'), 'UF'), motor.TramoD(date(2026, 1, 1), date(2026, 6, 30), D('25'), 'UF')]
        uf = lambda d: D('40000')
        con = motor.ContratoD(codigo='X', inicio=date(2024, 1, 1), tramos=t, renovacion_automatica=True)
        self.assertEqual(motor.renta_base(con, 2026, 11, uf).lineas[0].monto, D('25'))
        sin = motor.ContratoD(codigo='X', inicio=date(2024, 1, 1), tramos=t)
        self.assertTrue(any('Sin tramo' in a for a in motor.renta_base(sin, 2026, 11, uf).advertencias))


class ComparacionTest(Base):
    def _cargo(self, contrato, serv, mes, neto, cantidad=None):
        from .models import ConsumoServicio
        ConsumoServicio.objects.create(contrato=contrato, servicio=serv, periodo_consumo=date(2026, mes - 1, 1), periodo_cobro=date(2026, mes, 1),
                                       neto=neto, cantidad=cantidad, unidad_consumo='kWh' if cantidad else '')

    def test_semaforo_y_comparacion_con_3_meses(self):
        from . import comparacion
        self.assertEqual([comparacion.semaforo(D(x)) for x in ('0', '-10', '10.1', '20', '20.1', '-35')],
                         ['verde', 'verde', 'amarillo', 'amarillo', 'rojo', 'rojo'])
        for mes, kwh in ((8, 900), (9, 1000), (10, 1150)):
            self._cargo(self.fa, 'ELEC', mes, kwh * 150, D(kwh))
        self._cargo(self.fa, 'OTRO', 9, 100000)
        self._cargo(self.fa, 'OTRO', 10, 105000)
        self._cargo(self.nc, 'GAS', 10, 50000, D(40))
        meses, filas = comparacion.comparar(self.soc, date(2026, 10, 1))
        self.assertEqual(meses, [date(2026, 7, 1), date(2026, 8, 1), date(2026, 9, 1), date(2026, 10, 1)])
        f = {(x.contrato.codigo, x.servicio): x for x in filas}
        luz = f[('P-01', 'ELEC')]
        self.assertEqual((luz.variacion, luz.base, luz.luz), (D('15.0'), 'consumo', 'amarillo'))
        self.assertEqual((f[('P-01', 'OTRO')].variacion, f[('P-01', 'OTRO')].luz), (D('5.0'), 'verde'))
        self.assertEqual(f[('P-02', 'GAS')].luz, '')   # sin mes anterior: sin semáforo

    def test_pantalla_de_cargas_muestra_el_semaforo(self):
        self._cargo(self.fa, 'ELEC', 9, 100000, D(1000))
        self._cargo(self.fa, 'ELEC', 10, 130000, D(1300))
        Group.objects.create(name='Operación')
        u = User.objects.create_user('op', password='x')
        u.groups.add(Group.objects.get(name='Operación'))
        self.client.force_login(u)
        r = self.client.get(f'/cargas/?sociedad={self.soc.id}&periodo=2026-10')
        self.assertContains(r, '● rojo')
        self.assertContains(r, '+30 %')
        self.assertContains(self.client.get(f'/cargas/?sociedad={self.soc.id}&periodo=2026-10&alertas=1'), '1</b> en rojo')


class AsignacionIATest(Base):
    def test_la_ia_propone_y_lo_confirmado_se_asigna_solo_el_mes_siguiente(self):
        import os
        from unittest import mock
        from maestro import ia
        from maestro.tests import ClienteFalso
        from . import asignacion, lectores as L
        from .models import CargaPlanilla, ConsumoServicio
        filas = [L.FilaCargo('ELEC', '1221', 'Ale Cosmeticos', 'BT 1', D(1), D(246), None, D(245), 'kWh', D('55361'), 9),
                 L.FilaCargo('ELEC', 'AC', 'Area comun', 'BT 2', None, None, None, D(500), 'kWh', D('90000'), 10)]
        respuesta = {'asignaciones': [{'fila': 9, 'contrato': 'P-01', 'confianza': 'alta', 'motivo': 'Ale = nombre de fantasía de P-01'},
                                      {'fila': 10, 'contrato': 'NO', 'confianza': 'media', 'motivo': 'área común'}]}
        with mock.patch.dict(os.environ, {'ANTHROPIC_API_KEY': 'sk-ant-prueba'}), mock.patch.object(ia, 'EN_SEGUNDO_PLANO', False), \
                mock.patch.object(ia, '_cliente', return_value=ClienteFalso(respuesta)):
            carga = asignacion.preparar(self.soc, 'ELEC', filas, date(2026, 9, 1), date(2026, 10, 1), 'luz.xlsx', 'op')
        carga.refresh_from_db()
        self.assertEqual(carga.estado, 'L')
        self.assertEqual([(d['contrato'], d['fuente']) for d in carga.filas], [('P-01', 'IA'), ('NO', 'IA')])
        asignacion.confirmar(carga, {}, 'op')   # acepta lo propuesto
        self.assertEqual(ConsumoServicio.objects.get().contrato, self.fa)
        # el mes siguiente la misma fila se asigna sola por medidor confirmado, sin llamar a la IA
        carga2 = asignacion.preparar(self.soc, 'ELEC', filas[:1], date(2026, 10, 1), date(2026, 11, 1), 'luz2.xlsx', 'op')
        self.assertEqual((carga2.filas[0]['contrato'], carga2.filas[0]['fuente'], carga2.estado), ('P-01', 'medidor confirmado', 'L'))

    def test_las_pruebas_nunca_usan_la_clave_real(self):
        from maestro import ia
        self.assertEqual(ia.clave_api(), '')


class FechaUFTest(Base):
    def test_fecha_elegida_se_aplica_a_todos_incluso_primer_dia_habil(self):
        from .models import FechaUFMes
        UF.objects.create(fecha=date(2026, 10, 5), valor=D('41100.00'))
        self.nc.regla_uf = 'PH'
        self.nc.save()
        Group.objects.get_or_create(name='Operación')
        u = User.objects.create_user('op', password='x')
        u.groups.add(Group.objects.get(name='Operación'))
        self.client.force_login(u)
        r = self.client.post('/facturacion/fecha-uf/', {'sociedad': self.soc.id, 'periodo': '2026-10', 'fecha': '2026-10-09'}, follow=True)
        self.assertContains(r, 'no está cargada')   # sin UF de esa fecha no se acepta
        self.client.post('/facturacion/fecha-uf/', {'sociedad': self.soc.id, 'periodo': '2026-10', 'fecha': '2026-10-05'})
        self.assertEqual(FechaUFMes.objects.get().fecha, date(2026, 10, 5))
        servicios.generar_borradores(self.soc, 2026, 10)
        for d in Documento.objects.filter(periodo=date(2026, 10, 1)):
            self.assertEqual((d.fecha_uf, d.valor_uf), (date(2026, 10, 5), D('41100.00')))
            self.assertIn('05/10/2026', d.glosa)
        self.assertContains(self.client.get(f'/facturacion/?sociedad={self.soc.id}&periodo=2026-10'), 'elegida a mano')
        # noviembre no cambia
        servicios.generar_borradores(self.soc, 2026, 11)
        self.assertEqual(Documento.objects.get(contrato=self.fa, periodo=date(2026, 11, 1)).fecha_uf, date(2026, 11, 1))
        # volver al día 1
        self.client.post('/facturacion/fecha-uf/', {'sociedad': self.soc.id, 'periodo': '2026-10', 'accion': 'dia1'})
        self.assertFalse(FechaUFMes.objects.exists())
        servicios.generar_borradores(self.soc, 2026, 10)
        self.assertEqual(Documento.objects.get(contrato=self.fa, periodo=date(2026, 10, 1)).fecha_uf, date(2026, 10, 1))


class AguaYGastosComunesTest(Base):
    def _xlsx(self, filas):
        import openpyxl
        wb = openpyxl.Workbook()
        for f in filas:
            wb.active.append(f)
        d = tempfile.mkdtemp()
        ruta = Path(d) / 'p.xlsx'
        wb.save(ruta)
        return str(ruta)

    def test_lectura_flexible_de_planillas(self):
        from . import lectores
        ggcc = self._xlsx([['Gastos comunes octubre 2026'], [], ['Local', 'Arrendatario', 'm²', 'Gasto común neto'],
                           ['L01', 'Ale SpA', 24, 120500], ['L02', 'Casa Ltda', 30, 150000.4], ['', 'TOTAL', '', 270500]])
        hoja, filas = lectores.leer_generica(ggcc, 'GGCC')
        self.assertEqual([(f.etiqueta_local, f.nombre, f.neto) for f in filas], [('L01', 'Ale SpA', D('120500')), ('L02', 'Casa Ltda', D('150000.4'))])
        agua = self._xlsx([['N° Local', 'Nombre', 'Lectura anterior', 'Lectura actual', 'Consumo m3', 'Neto'], [1, 'Ale', 100, 112, 12, 15600]])
        _, filas = lectores.leer_generica(agua, 'AGUA')
        f = filas[0]
        self.assertEqual((f.etiqueta_local, f.cantidad, f.unidad, f.lectura_actual, f.neto), ('1', D('12'), 'm³', D('112'), D('15600')))
        with self.assertRaises(lectores.PlanillaInvalida):
            lectores.leer_generica(self._xlsx([['a', 'b'], [1, 2]]), 'AGUA')

    def test_gastos_comunes_en_la_factura_con_iva(self):
        from . import asignacion, lectores
        from .models import CargaPlanilla, LineaDocumento
        Contrato.objects.filter(pk__in=[self.fa.pk, self.nc.pk]).update(cobra_gastos_comunes='INM')   # los cobra la inmobiliaria
        ruta = self._xlsx([['Local', 'Arrendatario', 'Neto'], ['L01', 'Ale SpA', 100000], ['L02', 'Casa Ltda', 50000]])
        _, filas = lectores.leer_generica(ruta, 'GGCC')
        carga = asignacion.preparar(self.soc, 'GGCC', filas, date(2026, 9, 1), date(2026, 10, 1), 'ggcc.xlsx', 'op')
        asignacion.confirmar(carga, {2: 'P-01', 3: 'P-02'}, 'op')
        servicios.generar_borradores(self.soc, 2026, 10)
        fa = Documento.objects.get(contrato=self.fa)
        linea = fa.lineas.get(origen='GGCC')
        self.assertEqual((linea.concepto, int(linea.neto), linea.exento), ('GASTOS COMUNES', 100000, False))
        self.assertIn('Gastos comunes de Septiembre 2026', linea.detalle)
        self.assertNotIn('medidor', linea.detalle)   # gastos comunes: sin «medidor» (Matías, 2026-10-15)
        renta = sum(int(l.neto) for l in fa.lineas.exclude(origen='GGCC'))
        self.assertEqual(int(fa.iva), servicios.motor.peso(D(renta + 100000) * D('0.19')))   # IVA también sobre los gastos comunes
        nc = Documento.objects.get(contrato=self.nc)
        self.assertTrue(nc.lineas.get(origen='GGCC').exento)   # nota de cobro: sin IVA
        # el mes siguiente las filas confirmadas se asignan solas
        carga2 = asignacion.preparar(self.soc, 'GGCC', filas, date(2026, 10, 1), date(2026, 11, 1), 'ggcc2.xlsx', 'op')
        self.assertEqual([d['fuente'] for d in carga2.filas], ['medidor confirmado', 'medidor confirmado'])


    def test_los_que_cobra_megaservicios_no_van_en_la_factura(self):
        from . import asignacion, lectores
        from .models import ConsumoServicio
        Contrato.objects.filter(pk=self.fa.pk).update(cobra_agua='INM')   # P-01: agua la cobra la inmobiliaria; P-02 sigue con Megaservicios
        ruta = self._xlsx([['Local', 'Nombre', 'Neto'], ['L01', 'Ale', 8000], ['L02', 'Casa', 9000]])
        _, filas = lectores.leer_generica(ruta, 'AGUA')
        carga = asignacion.preparar(self.soc, 'AGUA', filas, date(2026, 9, 1), date(2026, 10, 1), 'agua.xlsx', 'op')
        asignacion.confirmar(carga, {2: 'P-01', 3: 'P-02'}, 'op')
        self.assertEqual(list(ConsumoServicio.objects.values_list('contrato__codigo', flat=True)), ['P-01'])
        self.assertIn('(los cobra Megaservicios)', carga.resultado)
        # si después se cambia P-01 a Megaservicios, lo ya cargado no entra en la factura
        Contrato.objects.filter(pk=self.fa.pk).update(cobra_agua='MEGA')
        servicios.generar_borradores(self.soc, 2026, 10)
        self.assertFalse(Documento.objects.get(contrato=self.fa).lineas.filter(origen='AGUA').exists())

    def test_pantalla_quien_cobra_cambia_y_deja_historial(self):
        from maestro.models import RegistroCambio
        Group.objects.get_or_create(name='Operación')
        u = User.objects.create_user('op', password='x')
        u.groups.add(Group.objects.get(name='Operación'))
        self.client.force_login(u)
        r = self.client.get(f'/contratos/quien-cobra/?sociedad={self.soc.id}')
        for txt in ('Megaservicios (directo al arrendatario)', '<th>Luz</th>', '<th>Gas</th>', '<th>Agua</th>', '<th>Gastos comunes</th>'):
            self.assertContains(r, txt)
        self.assertEqual((self.fa.cobra_luz, self.fa.cobra_gas), ('INM', 'INM'))   # por defecto luz y gas los cobra la inmobiliaria
        datos = {f'cobra_gastos_comunes_{self.fa.pk}': 'INM', f'cobra_agua_{self.fa.pk}': 'MEGA', f'cobra_luz_{self.fa.pk}': 'NO',
                 f'cobra_gas_{self.fa.pk}': 'INM', 'sociedad': self.soc.id}
        self.client.post(f'/contratos/quien-cobra/?sociedad={self.soc.id}', datos)
        self.fa.refresh_from_db()
        self.assertEqual((self.fa.cobra_gastos_comunes, self.fa.cobra_agua, self.fa.cobra_luz), ('INM', 'MEGA', 'NO'))
        self.assertFalse(self.fa.cobra_inmobiliaria('ELEC'))
        self.assertTrue(self.fa.cobra_inmobiliaria('GAS'))
        self.assertTrue(RegistroCambio.objects.filter(contrato=self.fa, usuario='op').exists())


class PlanillasPlazaMotorTest(Base):
    """Formatos de Plaza Motor (15 Oriente): agua con «Plaza Motor | Arrendatario | l. anterior m3 | … | $» y luz con
    columnas sin título de nombre, «Nro cliente», «Consumo Kwh» y «$»; hoja «AGO26»."""
    def _hoja(self, nombre, filas):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.title = nombre
        for f in filas:
            wb.active.append(f)
        ruta = Path(tempfile.mkdtemp()) / 'p.xlsx'
        wb.save(ruta)
        return str(ruta)

    def test_agua_plaza_motor(self):
        from . import lectores
        ruta = self._hoja('AGO26', [[], [None, None, 'Plaza Motor', 'Arrendatario', 'l. anterior m3', 'l. actual m3', 'consumo m3', '$'],
                                    [None, None, 1, 'Autoplanet', 1163, 1169, 6, 10359.214285714286],
                                    [None, None, 9, 'Totaland', 133, 135, 2, 3453.0714285714284]])
        hoja, filas = lectores.leer_generica(ruta, 'AGUA')
        self.assertEqual(lectores.mes_de_hoja(hoja), (2026, 8))
        self.assertEqual([(f.etiqueta_local, f.nombre, f.lectura_anterior, f.cantidad, f.unidad) for f in filas],
                         [('1', 'Autoplanet', D('1163'), D('6'), 'm³'), ('9', 'Totaland', D('133'), D('2'), 'm³')])   # «Totaland» no es una fila de total

    def test_luz_plaza_motor_sin_titulo_de_nombre(self):
        from . import lectores
        ruta = self._hoja('AGO26', [[], [], [None, None, 'Consumos Electricidad', None, 'Nro cliente', 'Consumo Kwh', '$'],
                                    [None, None, 1, 'Autoplanet', None, None, None],
                                    [None, None, 2, 'Würth', 6349302, 541, 118061],
                                    [None, None, 9, ' +56 Food', 6248866, None, None],
                                    [None, None, 10, ' +56 Market', 6252582, 1267, 412913]])
        _, filas = lectores.leer_generica(ruta, 'ELEC')
        self.assertEqual([(f.etiqueta_local, f.nombre, f.tarifa, f.cantidad, f.unidad, f.neto) for f in filas],
                         [('2', 'Würth', 'cliente 6349302', D('541'), 'kWh', D('118061')),
                          ('10', '+56 Market', 'cliente 6252582', D('1267'), 'kWh', D('412913'))])   # sin monto: no se carga

    def test_mes_de_hoja(self):
        from . import lectores
        self.assertEqual(lectores.mes_de_hoja('SEP26'), (2026, 9))
        self.assertEqual(lectores.mes_de_hoja('Octubre 2026'), (2026, 10))
        self.assertIsNone(lectores.mes_de_hoja('Hoja1'))

    def test_subir_luz_plaza_motor_y_avisar_mes_distinto(self):
        Group.objects.get_or_create(name='Operación')
        u = User.objects.create_user('op2', password='x')
        u.groups.add(Group.objects.get(name='Operación'))
        self.client.force_login(u)
        ruta = self._hoja('AGO26', [[None, 'Consumos Electricidad', None, 'Nro cliente', 'Consumo Kwh', '$'],
                                    [None, 1, 'Ale', 111, 10, 5000], [None, 2, 'Casa', 222, 20, 9000]])
        def subir(consumo):
            with open(ruta, 'rb') as fh:
                return self.client.post(f'/cargas/archivo/?sociedad={self.soc.id}&periodo=2026-09',
                                        {'tipo': 'electricidad', 'consumo': consumo, 'archivo': fh}, follow=True)
        self.assertContains(subir('2026-09'), 'es de Agosto 2026')   # la hoja dice agosto y se indicó septiembre
        r = subir('2026-08')
        self.assertEqual(r.redirect_chain[-1][0].split('/')[1], 'cargas')
        from .models import CargaPlanilla
        carga = CargaPlanilla.objects.get()
        self.assertEqual((carga.servicio, len(carga.filas)), ('ELEC', 2))


class DetalleXMLTest(Base):
    def test_cada_linea_lleva_el_texto_del_calculo(self):
        docs = self.generar_y_aprobar()
        fa = docs['P-01']
        xml = bsoft.xml_documento(fa, {o: '1002' for o in ('RENTA', 'RDIF', 'DESC', 'ELEC', 'GAS', 'OTRO', 'AGUA', 'GGCC')})
        linea = fa.lineas.order_by('orden').first()
        self.assertIn(f'<DscItem>{bsoft.detalle_linea(linea)}</DscItem>', xml)
        self.assertTrue(bsoft.detalle_linea(linea).startswith(linea.concepto))
        self.assertIn('41.065,38', bsoft.detalle_linea(linea))   # el cálculo con la UF usada
        self.assertTrue(xml.isascii())

    def test_texto_sin_simbolos_y_con_largo_maximo(self):
        from types import SimpleNamespace
        l = SimpleNamespace(concepto='ARRIENDO', detalle='27 UF × UF 41.065,38 → año ' + 'x' * 2000)
        t = bsoft.detalle_linea(l)
        self.assertTrue(t.isascii())
        self.assertIn('27 UF x UF 41.065,38 -> ano', t)
        self.assertEqual(len(t), bsoft.MAX_DETALLE)


class GastosComunesAnualTest(Base):
    _hoja = PlanillasPlazaMotorTest._hoja

    def test_toma_la_columna_del_mes(self):
        from . import lectores
        meses = ['ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO', 'AGOSTO', 'SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE']
        ruta = self._hoja('GGCC 2026', [[None, 'GASTOS COMUNES PLAZA MOTOR'], [None, None, 2026, None, None, None] + meses,
                                        [None, None, 'TOTAL GGCC', None, None, None] + [1000] * 12,
                                        [None, 'Local', 'Arrendatario', 'm2', '%', 'Max UF/m2'] + meses,
                                        [None, 1, 'Autoplanet', 478, 0.37, 0.1] + [100.4] * 7 + [1379807.79, 505125.6, None, None, None],
                                        [None, '4 y 5', 'Disponible', 378, 0.29, None] + [None] * 12,
                                        [None, 9, 'Totaland', 50.74, 0.04, None] + [5] * 7 + [137775.1, None, None, None, None],
                                        [None, 'Total', None, 1291.44, 1, None]])
        hoja, filas = lectores.leer_generica(ruta, 'GGCC', 8)
        self.assertEqual(lectores.anio_de_hoja(hoja), 2026)
        self.assertEqual([(f.etiqueta_local, f.nombre, f.neto) for f in filas],
                         [('1', 'Autoplanet', D('1379807.79')), ('9', 'Totaland', D('137775.1'))])
        _, filas = lectores.leer_generica(ruta, 'GGCC', 9)
        self.assertEqual([f.nombre for f in filas], ['Autoplanet'])   # Totaland sin monto en septiembre


class LargoCamposTest(TestCase):
    def test_numero_de_cliente_cabe_en_tarifa(self):
        """PostgreSQL (nube) rechaza textos más largos que el campo; «cliente 6349302» debe caber (error 500 del 2026-10-14)."""
        from .models import ConsumoServicio, MedidorServicio
        for m in (ConsumoServicio, MedidorServicio):
            self.assertGreaterEqual(m._meta.get_field('tarifa').max_length, len('cliente 12345678901'))


class TextosLargosTest(Base):
    """Ningún texto demasiado largo debe botar el sistema en la nube (PostgreSQL): se recorta al guardar."""
    def test_en_pruebas_avisa_y_en_el_sistema_recorta(self):
        from maestro import largos
        from .models import MedidorServicio
        m = MedidorServicio(contrato=self.fa, servicio='ELEC', nombre_planilla='x' * 200, tarifa='t')
        with self.assertRaises(largos.TextoMuyLargo):
            m.save()
        with self.settings(PRUEBAS=False):
            m.save()
        m.refresh_from_db()
        self.assertEqual(len(m.nombre_planilla), 80)

    def test_planilla_con_nombres_y_textos_muy_largos_se_carga(self):
        from . import asignacion, lectores
        from .models import ConsumoServicio
        Contrato.objects.filter(pk=self.fa.pk).update(cobra_agua='INM')
        largo = 'Arrendatario con un nombre comercial larguísimo ' * 6
        ruta = AguaYGastosComunesTest._xlsx(self, [['N° local de la planilla del proveedor ' * 3, 'Nombre', 'Nro cliente', 'Consumo m3', 'Neto'],
                                                   ['L' * 90, largo, 123456789012345678901234567890123456789, 5, 7000]])
        _, filas = lectores.leer_generica(ruta, 'AGUA')
        with self.settings(PRUEBAS=False):
            carga = asignacion.preparar(self.soc, 'AGUA', filas, date(2026, 9, 1), date(2026, 10, 1), 'x' * 300 + '.xlsx', 'op')
            asignacion.confirmar(carga, {2: 'P-01'}, 'op')
        self.assertEqual(ConsumoServicio.objects.get().neto, 7000)


class TramosSuperpuestosTest(Base):
    def test_dos_tramos_a_la_vez_dejan_advertencia_y_no_se_aprueba(self):
        """Caso O-01 (2026-10-14): los cuatro tramos quedaron con el mismo «desde» y el sistema cobraba uno sin avisar."""
        TramoRenta.objects.create(contrato=self.fa, desde=date(2025, 1, 1), monto=D('30'), unidad='UF')
        servicios.generar_borradores(self.soc, 2026, 10)
        doc = Documento.objects.get(contrato=self.fa)
        self.assertTrue(any('tramos de renta vigentes a la vez' in a and '27' in a and '30' in a for a in doc.advertencias))
        servicios.aprobar(self.soc, 2026, 10, 'op')
        doc.refresh_from_db()
        self.assertEqual(doc.estado, 'B')   # queda en borrador hasta corregir las fechas

    def test_tramos_consecutivos_no_avisan(self):
        TramoRenta.objects.filter(contrato=self.fa).update(hasta=date(2026, 9, 30))
        TramoRenta.objects.create(contrato=self.fa, desde=date(2026, 10, 1), monto=D('30'), unidad='UF')
        servicios.generar_borradores(self.soc, 2026, 10)
        doc = Documento.objects.get(contrato=self.fa)
        self.assertFalse(any('a la vez' in a for a in doc.advertencias))
        self.assertIn('30', doc.lineas.get(origen='RENTA').detalle)


class ExcelFacturacionTest(Base):
    """Excel de facturación (Matías, 2026-10-15): glosa y detalle juntos sin tildes ni ñ, código de artículo y condición de pago."""
    def test_columnas_nuevas_y_texto_sin_tildes(self):
        import openpyxl
        from . import exportar
        from .models import ConsumoServicio
        from maestro.models import CondicionPago
        Contrato.objects.filter(pk=self.nc.pk).update(condicion_pago=CondicionPago.objects.get(nombre='CONTADO'))
        ConsumoServicio.objects.create(contrato=self.fa, servicio='OTRO', concepto='Habilitación del local', detalle='Diseño señalética',
                                       periodo_consumo=date(2026, 10, 1), periodo_cobro=date(2026, 10, 1), neto=50000)
        CodigoServicioBsoft.objects.filter(sociedad=self.soc).delete()   # sin tabla: usa los códigos estándar
        self.generar_y_aprobar()
        ruta = exportar.exportar_excel(self.soc, 2026, 10, ruta=str(Path(tempfile.mkdtemp()) / 'f.xlsx'))
        ws = openpyxl.load_workbook(ruta)['Facturación']
        enc = [c.value for c in ws[5]]
        self.assertNotIn('Glosa del documento', enc)
        self.assertNotIn('Detalle del cálculo', enc)
        filas = [dict(zip(enc, [c.value for c in f])) for f in ws.iter_rows(min_row=6) if f[1].value]
        por = {(f['Contrato'], f['Concepto']): f for f in filas}
        renta = next(f for f in filas if f['Contrato'] == 'P-01' and f['Código artículo'] == '111')
        self.assertTrue(renta['Glosa y detalle del cálculo'].startswith('ARRIENDO - Octubre 2026 - Calculo: '))
        self.assertEqual(renta['Glosa y detalle del cálculo'].count('41.065,38'), 1)   # la UF aparece una sola vez
        self.assertIn('Calculo: 27,00 UF', renta['Glosa y detalle del cálculo'])
        self.assertEqual(renta['Condición de pago'], 'CREDITO')   # por defecto
        self.assertEqual(renta['N° interno'], 'IP-F-0001')
        otro = next(f for f in filas if f['Contrato'] == 'P-01' and 'HABILITACION' in (f['Concepto'] or ''))
        self.assertEqual(otro['Código artículo'], '222')
        self.assertIn('Diseno senaletica', otro['Glosa y detalle del cálculo'])
        nc = next(f for f in filas if f['Contrato'] == 'P-02')
        self.assertEqual(nc['Condición de pago'], 'CONTADO')
        for f in filas:
            self.assertTrue(all(str(v).isascii() for k, v in f.items() if k in ('Glosa y detalle del cálculo', 'Concepto', 'Condición de pago')))

    def test_codigos_de_articulo(self):
        from types import SimpleNamespace as N
        from . import articulos
        cod = {'RENTA': '111', 'GGCC': '777'}
        self.assertEqual(articulos.codigo(N(origen='GGCC', concepto='GASTOS COMUNES', detalle=''), cod), '777')
        self.assertEqual(articulos.codigo(N(origen='DESC', concepto='DESCUENTO', detalle=''), {}), '111')
        self.assertEqual(articulos.codigo(N(origen='OTRO', concepto='Gastos administrativos', detalle=''), cod), '777')
        self.assertEqual(articulos.codigo(N(origen='OTRO', concepto='Arriendo de equipamiento', detalle=''), cod), '888')
        self.assertEqual(articulos.codigo(N(origen='OTRO', concepto='Reembolso', detalle=''), cod), '999')
        self.assertEqual(articulos.codigo(N(origen='AGUA', concepto='AGUA', detalle=''), {}), '444')


class PantallaCodigosTest(Base):
    def test_tabla_por_empresa_editable_solo_por_admin(self):
        Group.objects.get_or_create(name='Operación')
        op = User.objects.create_user('op3', password='x')
        op.groups.add(Group.objects.get(name='Operación'))
        self.client.force_login(op)
        url = f'/facturacion/codigos/?sociedad={self.soc.id}'
        r = self.client.get(url)
        self.assertContains(r, 'Gastos comunes')
        self.assertContains(r, 'value="1002"')   # código actual de la renta en esta sociedad de prueba
        self.assertEqual(self.client.post(url, {'c_RENTA': '111'}).status_code, 403)
        self.client.force_login(User.objects.create_superuser('jefe2', password='x'))
        self.client.post(url, {'c_RENTA': '111', 'd_RENTA': 'ARRIENDO', 'c_GGCC': '777', 'd_GGCC': 'GASTOS ADMINISTRATIVOS'})
        cod = dict(CodigoServicioBsoft.objects.filter(sociedad=self.soc).values_list('origen', 'codigo'))
        self.assertEqual((cod['RENTA'], cod['GGCC'], cod['ELEC']), ('111', '777', '1002'))



class CondicionPagoTest(Base):
    """Condiciones de pago en el maestro (Matías, 2026-10-15): lista desplegable en el contrato, CRÉDITO por defecto,
    y se muestra en el documento, en el PDF de la nota de cobro y en el Excel."""
    def test_maestro_por_defecto_y_en_la_factura(self):
        from maestro.models import CondicionPago
        self.assertEqual(list(CondicionPago.objects.values_list('nombre', flat=True)), ['CONTADO', 'CRÉDITO', 'CRÉDITO 30 DÍAS', 'CRÉDITO 60 DÍAS'])
        self.fa.refresh_from_db()
        self.assertEqual(str(self.fa.condicion_pago), 'CRÉDITO')   # contrato nuevo: CRÉDITO
        c60 = CondicionPago.objects.get(nombre='CRÉDITO 60 DÍAS')
        Contrato.objects.filter(pk=self.nc.pk).update(condicion_pago=c60)
        docs = self.generar_y_aprobar()
        self.assertEqual((docs['P-01'].condicion_pago, docs['P-02'].condicion_pago), ('CRÉDITO', 'CRÉDITO 60 DÍAS'))
        texto = ''.join(p.extract_text() for p in PdfReader(io.BytesIO(nota_cobro.pdf_nota(docs['P-02']))).pages)
        self.assertIn('60 D', texto)
        Group.objects.get_or_create(name='Operación')
        u = User.objects.create_user('op4', password='x')
        u.groups.add(Group.objects.get(name='Operación'))
        self.client.force_login(u)
        self.assertContains(self.client.get(f'/documento/{docs["P-02"].pk}/'), 'CRÉDITO 60 DÍAS')
        r = self.client.get('/contratos/P-01/editar/')
        self.assertContains(r, '<select name="condicion_pago"')
        self.assertNotContains(r, '---------')   # sin opción vacía



class LuzNoSeCobraTest(Base):
    def test_luz_que_no_se_cobra_no_va_en_la_factura(self):
        from .models import ConsumoServicio
        Contrato.objects.filter(pk=self.fa.pk).update(cobra_luz='NO')
        for c in (self.fa, self.nc):
            ConsumoServicio.objects.create(contrato=c, servicio='ELEC', periodo_consumo=date(2026, 9, 1), periodo_cobro=date(2026, 10, 1), neto=10000)
        servicios.generar_borradores(self.soc, 2026, 10)
        fa, nc = Documento.objects.get(contrato=self.fa), Documento.objects.get(contrato=self.nc)
        self.assertFalse(fa.lineas.filter(origen='ELEC').exists())
        self.assertTrue(nc.lineas.filter(origen='ELEC').exists())


class AnexoDeRentaTest(Base):
    def test_anexo_con_fecha_posterior_manda_sin_advertencia(self):
        """Anexo O-01: «se acuerda fijar renta desde enero 2026 a diciembre 2026 en UF 187,65» sobre tramos antiguos."""
        TramoRenta.objects.create(contrato=self.fa, desde=date(2025, 1, 1), monto=D('30'), unidad='UF')   # dos tramos antiguos empatados
        TramoRenta.objects.create(contrato=self.fa, desde=date(2026, 1, 1), hasta=date(2026, 12, 31), monto=D('187.65'), unidad='UF')
        servicios.generar_borradores(self.soc, 2026, 10)
        doc = Documento.objects.get(contrato=self.fa)
        self.assertFalse(any('a la vez' in a for a in doc.advertencias))
        self.assertIn('187,65 UF', doc.lineas.get(origen='RENTA').detalle)
