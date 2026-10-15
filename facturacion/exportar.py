"""Excel de facturación genérico (una línea por concepto, por sociedad). El formato exacto de importación
de BSOFT es la etapa 2: cuando llegue la plantilla solo cambia este módulo."""
import os
from datetime import date

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

from . import articulos, motor
from .models import CodigoServicioBsoft, Documento

FN = 'Arial'
f_norm = Font(name=FN, size=10)
f_bold = Font(name=FN, size=10, bold=True)
f_input = Font(name=FN, size=10, color='0000FF')
f_head = Font(name=FN, size=10, bold=True, color='FFFFFF')
f_title = Font(name=FN, size=14, bold=True, color='1F3A5F')
f_sub = Font(name=FN, size=10, italic=True, color='595959')
fill_head = PatternFill('solid', start_color='1F3A5F')
fill_warn = PatternFill('solid', start_color='FFE699')
thin = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')


def _head(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.border = f_head, fill_head, border
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[CL(i)].width = w
    ws.row_dimensions[row].height = 34


def numero(d) -> str:
    """N.º interno del documento (BT-F-0001…); los borradores aún no lo tienen."""
    return d.numero_interno or f'Borrador {d.id}'


def glosa_y_detalle(d, l) -> str:
    """Glosa del documento y detalle del cálculo de la línea en un solo texto, sin tildes ni ñ:
    «ARRIENDO - Octubre 2026 - UF 41.098,15 del 05/10/2026 - Calculo: 202,5 UF mes completo x UF 41.098,15»."""
    partes = [l.concepto, d.glosa.replace('Valor de la UF', 'UF')]
    if l.detalle:
        partes.append('Calculo: ' + l.detalle)
    return articulos.sin_tildes(' - '.join(p.strip() for p in partes if p and p.strip()))


def exportar_excel(sociedad, año: int, mes: int, estados=('A', 'X'), fecha_emision: date = None, ruta=None) -> str:
    docs = list(Documento.objects.filter(contrato__sociedad=sociedad, periodo=date(año, mes, 1), estado__in=estados)
                .select_related('contrato__arrendatario', 'contrato__propiedad').prefetch_related('lineas', 'contrato__locales').order_by('contrato__codigo'))
    fecha_emision = fecha_emision or date(año, mes, 1)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Facturación'
    ws.sheet_view.showGridLines = False
    ws['A1'] = f'Facturación {motor.MESES[mes]} {año} · {sociedad.nombre}'
    ws['A1'].font = f_title
    ws['A2'] = f'RUT {sociedad.rut} · una línea por concepto · estados incluidos: {", ".join(dict(Documento.Estado.choices)[e] for e in estados)}'
    ws['A2'].font = f_sub
    ws['A3'], ws['B3'] = 'Tasa de IVA', 0.19
    ws['A3'].font, ws['B3'].font = f_bold, f_input
    ws['B3'].number_format = '0%'
    H = ['N° interno', 'Contrato', 'Tipo de documento', 'Fecha de emisión', 'RUT cliente', 'Razón social', 'Nombre de fantasía', 'Dirección (según contrato)',
         'Condición de pago', 'Línea', 'Código artículo', 'Concepto', 'Glosa y detalle del cálculo', 'Unidad', 'Monto en UF o $', 'Valor UF',
         'Neto $', 'Exento (1 = sí)', 'IVA $', 'Total $']
    W = [12, 9, 14, 12, 14, 34, 24, 30, 12, 6, 10, 26, 80, 7, 14, 11, 13, 9, 12, 13]
    col = {h: n for n, h in enumerate(H, start=1)}
    L = lambda h: CL(col[h])
    _head(ws, 5, H, W)
    r = 6
    tipo = dict(Documento.Estado.choices)
    tipos = {'FA': 'Factura afecta', 'FE': 'Factura exenta', 'NC': 'Nota de cobro'}
    codigos = dict(CodigoServicioBsoft.objects.filter(sociedad=sociedad).values_list('origen', 'codigo'))
    for d in docs:
        c = d.contrato
        for l in d.lineas.all():
            vals = {'N° interno': numero(d), 'Contrato': c.codigo, 'Tipo de documento': tipos.get(d.tipo_documento, d.tipo_documento),
                    'Fecha de emisión': fecha_emision, 'RUT cliente': c.arrendatario.rut, 'Razón social': c.arrendatario.razon_social,
                    'Nombre de fantasía': c.nombre_fantasia, 'Dirección (según contrato)': c.direccion_contrato,
                    'Condición de pago': articulos.sin_tildes(c.condicion_pago or 'CREDITO').upper(), 'Línea': l.orden,
                    'Código artículo': articulos.codigo(l, codigos), 'Concepto': articulos.sin_tildes(l.concepto),
                    'Glosa y detalle del cálculo': glosa_y_detalle(d, l), 'Unidad': l.unidad, 'Monto en UF o $': float(l.cantidad_unidad),
                    'Valor UF': float(l.valor_uf) if l.valor_uf else None, 'Neto $': int(l.neto), 'Exento (1 = sí)': 1 if l.exento else 0}
            for h, v in vals.items():
                cell = ws.cell(row=r, column=col[h], value=v)
                cell.font = f_input if h in ('Monto en UF o $', 'Valor UF', 'Neto $') else f_norm
                cell.border, cell.alignment = border, wrap
            ws.cell(row=r, column=col['Fecha de emisión']).number_format = 'dd-mm-yyyy'
            ws.cell(row=r, column=col['Monto en UF o $']).number_format = '#,##0.0000' if l.unidad == 'UF' else '#,##0'
            ws.cell(row=r, column=col['Valor UF']).number_format = '#,##0.00'
            ws.cell(row=r, column=col['Neto $']).number_format = '#,##0;-#,##0'
            f_iva = ws.cell(row=r, column=col['IVA $'], value=f"=IF({L('Exento (1 = sí)')}{r}=1,0,ROUND({L('Neto $')}{r}*$B$3,0))")
            f_tot = ws.cell(row=r, column=col['Total $'], value=f"={L('Neto $')}{r}+{L('IVA $')}{r}")
            for cc in (f_iva, f_tot):
                cc.font, cc.border, cc.number_format = f_norm, border, '#,##0;-#,##0'
            r += 1
    ultimo = r - 1
    ws.cell(row=r + 1, column=col['Concepto'], value='Totales').font = f_bold
    for k in (col['Neto $'], col['IVA $'], col['Total $']):
        c_ = ws.cell(row=r + 1, column=k, value=f'=SUM({CL(k)}6:{CL(k)}{max(ultimo, 6)})')
        c_.font, c_.number_format = f_bold, '#,##0;-#,##0'
    ws.freeze_panes = 'F6'
    ws.auto_filter.ref = f'A5:{CL(len(H))}{max(ultimo, 6)}'

    rs = wb.create_sheet('Resumen')
    rs.sheet_view.showGridLines = False
    rs['A1'] = 'Resumen por documento'
    rs['A1'].font = f_title
    rs['A2'] = 'Valores del motor de cálculo (IVA calculado sobre el neto del documento). La última columna compara con la suma de IVA de las líneas de la hoja Facturación.'
    rs['A2'].font = f_sub
    HR = ['N° interno', 'Contrato', 'Fantasía', 'Locales', 'Estado', 'Neto $', 'Exento $', 'IVA $', 'Total $', 'Advertencias', 'IVA líneas (hoja Facturación)', 'Diferencia IVA']
    _head(rs, 4, HR, [10, 9, 28, 20, 11, 13, 12, 12, 13, 50, 16, 12])
    for i, d in enumerate(docs):
        rr = 5 + i
        locs = ', '.join(l.codigo_plano for l in d.contrato.locales.all())
        row = [numero(d), d.contrato.codigo, d.contrato.nombre_fantasia, locs, tipo[d.estado], int(d.neto), int(d.exento), int(d.iva), int(d.total), '; '.join(d.advertencias) or None]
        for j, v in enumerate(row, start=1):
            c_ = rs.cell(row=rr, column=j, value=v)
            c_.font, c_.border, c_.alignment = (f_input if j in (6, 7, 8, 9) else f_norm), border, wrap
            if j in (6, 7, 8, 9):
                c_.number_format = '#,##0'
        c_ = rs.cell(row=rr, column=11, value=f"=SUMIFS(Facturación!${L('IVA $')}$6:${L('IVA $')}${max(ultimo, 6)},Facturación!$A$6:$A${max(ultimo, 6)},A{rr})")
        c_.font, c_.border, c_.number_format = f_norm, border, '#,##0'
        c_ = rs.cell(row=rr, column=12, value=f'=H{rr}-K{rr}')
        c_.font, c_.border, c_.number_format = f_norm, border, '#,##0;-#,##0;0'
        if d.advertencias:
            rs.cell(row=rr, column=10).fill = fill_warn
    fin = 4 + len(docs)
    rs.cell(row=fin + 2, column=3, value='Totales').font = f_bold
    for col in (6, 7, 8, 9, 11, 12):
        c_ = rs.cell(row=fin + 2, column=col, value=f'=SUM({CL(col)}5:{CL(col)}{max(fin, 5)})')
        c_.font, c_.number_format = f_bold, '#,##0;-#,##0;0'
    rs.freeze_panes = 'D5'

    ruta = ruta or os.path.join(settings.MEDIA_ROOT, 'facturacion', f'Facturacion_{sociedad.rut.replace(".", "")}_{año}-{mes:02d}.xlsx')
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    wb.save(ruta)
    return ruta
