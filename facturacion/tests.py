"""Pruebas del motor de rentas (puras) y del flujo borrador → aprobación → Excel/XML (con base de datos)."""
import tempfile
from datetime import date
from decimal import Decimal as D

from django.test import SimpleTestCase, TestCase, override_settings

from facturacion import bsoft, motor, servicios
from facturacion.models import CodigoServicioBsoft, Documento, VentaMensual
from indicadores.models import UF
from maestro.models import Arrendatario, Contrato, Propiedad, Sociedad, TramoRenta

T = motor.TramoD
UFS = {date(2026, 9, 1): D('40875.09'), date(2026, 10, 1): D('41065.38'), date(2026, 11, 1): D('41200.00'),
       date(2026, 12, 1): D('41300.00'), date(2027, 1, 1): D('41400.00')}


def uf(d):
    try:
        return UFS[d]
    except KeyError:
        raise motor.UFNoDisponible(d)


def contrato(**kw):
    base = dict(codigo='X-01', inicio=date(2025, 1, 1))
    base.update(kw)
    return motor.ContratoD(**base)


class MotorTests(SimpleTestCase):
    def test_renta_fija_mes_completo(self):
        c = contrato(tramos=[T(date(2025, 1, 1), None, D('27'), 'UF')])
        r = motor.calcular_mes(c, 2026, 10, uf)
        self.assertEqual(r.neto, 1108765)  # 27 × 41.065,38 = 1.108.765,26
        self.assertEqual(r.valor_uf, D('41065.38'))
        self.assertEqual(r.fecha_uf, date(2026, 10, 1))

    def test_redondeo_mitad_hacia_arriba(self):
        self.assertEqual(motor.peso(D('0.5')), 1)
        self.assertEqual(motor.peso(D('2.5')), 3)
        self.assertEqual(motor.peso(D('2.4999')), 2)

    def test_gracia_no_genera_linea(self):
        c = contrato(inicio=date(2026, 9, 14), tramos=[T(date(2026, 9, 1), date(2026, 10, 31), D('0'), 'UF'),
                                                      T(date(2026, 11, 1), None, D('40'), 'UF')])
        r = motor.calcular_mes(c, 2026, 10, uf)
        self.assertEqual(r.lineas, [])
        self.assertTrue(r.notas)
        self.assertEqual(motor.calcular_mes(c, 2026, 11, uf).neto, peso_uf(40, '41200.00'))

    def test_prorrata_cambio_a_mitad_de_mes(self):
        # B-13: 42 UF hasta el 06-10 y 44 UF desde el 07-10 → 6/31 y 25/31
        c = contrato(tramos=[T(date(2025, 10, 7), date(2026, 10, 6), D('42'), 'UF'), T(date(2026, 10, 7), None, D('44'), 'UF')])
        r = motor.calcular_mes(c, 2026, 10, uf)
        esperado = motor.peso((D('42') * 6 / 31 + D('44') * 25 / 31) * D('41065.38'))
        self.assertEqual(r.neto, esperado)
        self.assertEqual(r.neto, 1790980)

    def test_modo_inicio_de_mes_no_prorratea(self):
        c = contrato(modo_cambio='IM', tramos=[T(date(2025, 10, 7), date(2026, 10, 6), D('42'), 'UF'), T(date(2026, 10, 7), None, D('44'), 'UF')])
        self.assertEqual(motor.calcular_mes(c, 2026, 10, uf).neto, peso_uf(42, '41065.38'))

    def test_primer_mes_proporcional(self):
        c = contrato(inicio=date(2026, 10, 16), tramos=[T(date(2026, 10, 16), None, D('31'), 'UF')])
        r = motor.calcular_mes(c, 2026, 10, uf)  # 16 días de 31
        self.assertEqual(r.neto, motor.peso(D('31') * 16 / 31 * D('41065.38')))

    def test_contrato_terminado_no_factura(self):
        c = contrato(fin=date(2026, 9, 30), tramos=[T(date(2025, 1, 1), None, D('10'), 'UF')])
        self.assertEqual(motor.calcular_mes(c, 2026, 10, uf).lineas, [])

    def test_uf_no_disponible_es_error_claro(self):
        c = contrato(tramos=[T(date(2025, 1, 1), None, D('10'), 'UF')])
        r = motor.calcular_mes(c, 2027, 6, uf)  # no inventa un valor: deja el mes sin línea y avisa
        self.assertEqual(r.lineas, [])
        self.assertTrue(any('UF' in a for a in r.advertencias))

    def test_renta_clp_no_usa_uf(self):
        c = contrato(tramos=[T(date(2025, 1, 1), None, D('2000000'), 'CLP')])
        self.assertEqual(motor.calcular_mes(c, 2026, 10, uf).neto, 2000000)

    def test_variable_agrega_diferencia(self):
        # mínimo 333,33 UF; 2% de $900.000.000 = $18.000.000 > mínimo de septiembre (333,33 × 40.875,09)
        c = contrato(tramos=[T(date(2025, 1, 1), None, D('333.33'), 'UF')], pct_variable=D('0.02'), base_variable='V')
        r = motor.calcular_mes(c, 2026, 10, uf, ventas_ant=D('900000000'))
        minimo_sep = motor.peso(D('333.33') * D('40875.09'))
        dif = 18000000 - minimo_sep
        self.assertEqual([l.origen for l in r.lineas], ['RENTA', 'RDIF'])
        self.assertEqual(r.lineas[1].neto, dif)

    def test_variable_bajo_el_minimo_no_cobra_diferencia(self):
        c = contrato(tramos=[T(date(2025, 1, 1), None, D('333.33'), 'UF')], pct_variable=D('0.02'), base_variable='V')
        r = motor.calcular_mes(c, 2026, 10, uf, ventas_ant=D('100000000'))
        self.assertEqual([l.origen for l in r.lineas], ['RENTA'])

    def test_variable_sin_ventas_advierte(self):
        c = contrato(tramos=[T(date(2025, 1, 1), None, D('333.33'), 'UF')], pct_variable=D('0.02'), base_variable='V')
        r = motor.calcular_mes(c, 2026, 10, uf)
        self.assertEqual(len(r.lineas), 1)
        self.assertTrue(any('Faltan las ventas' in a for a in r.advertencias))

    def test_variable_desde_posterior_ignora_meses_anteriores(self):
        c = contrato(tramos=[T(date(2025, 1, 1), None, D('10'), 'CLP')], pct_variable=D('0.5'), base_variable='I', variable_desde=date(2027, 1, 1))
        r = motor.calcular_mes(c, 2027, 1, uf, ventas_ant=D('99999999'))  # el «mes anterior» (dic-2026) aún no tenía variable
        self.assertEqual([l.origen for l in r.lineas], ['RENTA'])
        self.assertEqual(r.advertencias, [])

    def test_semestral_solo_en_meses_de_cobro(self):
        c = contrato(inicio=date(2026, 2, 1), periodicidad='S', meses_cobro=[1, 6], reajuste_ipc=True,
                     tramos=[T(date(2026, 2, 1), None, D('2500000'), 'CLP')])
        ipc = lambda a, m: D('0.5')
        self.assertEqual(motor.calcular_mes(c, 2026, 10, uf, ipc_fn=ipc).lineas, [])
        self.assertEqual(motor.calcular_mes(c, 2026, 6, uf, ipc_fn=ipc).neto, 2500000)

    def test_semestral_ipc_enero(self):
        c = contrato(inicio=date(2026, 2, 1), periodicidad='S', meses_cobro=[1, 6], reajuste_ipc=True,
                     tramos=[T(date(2026, 2, 1), None, D('2500000'), 'CLP')])
        ipc = lambda a, m: D('1')  # 1 % mensual feb–dic 2026 = 11 meses
        r = motor.calcular_mes(c, 2027, 1, uf, ipc_fn=ipc)
        self.assertEqual(r.neto, motor.peso(D('2500000') * (D('1.01') ** 11)))

    def test_semestral_ipc_pendiente_advierte(self):
        c = contrato(inicio=date(2026, 2, 1), periodicidad='S', meses_cobro=[1, 6], reajuste_ipc=True,
                     tramos=[T(date(2026, 2, 1), None, D('2500000'), 'CLP')])
        r = motor.calcular_mes(c, 2027, 1, uf, ipc_fn=lambda a, m: None)
        self.assertEqual(r.neto, 2500000)
        self.assertTrue(any('IPC pendiente' in a for a in r.advertencias))

    def test_descuento_porcentual_prorrateado(self):
        c = contrato(tramos=[T(date(2025, 1, 1), None, D('1000000'), 'CLP')],
                     descuentos=[motor.DescuentoD(date(2026, 10, 1), date(2026, 10, 10), 'PCT', D('50'), 'convenio')])
        r = motor.calcular_mes(c, 2026, 10, uf)
        self.assertEqual(r.lineas[1].neto, -motor.peso(D(1000000) * D('0.5') * 10 / 31))

    def test_regla_primer_dia_habil(self):
        c = contrato(regla_uf='PH', tramos=[T(date(2025, 1, 1), None, D('1'), 'UF')])
        # 1-nov-2026 es domingo → primer día hábil lunes 2-nov (se pasa un feriado de ejemplo el 2)
        self.assertEqual(motor.fecha_uf_del_mes(c, 2026, 11), date(2026, 11, 2))
        self.assertEqual(motor.fecha_uf_del_mes(c, 2026, 11, feriados=(date(2026, 11, 2),)), date(2026, 11, 3))

    def test_totales_iva(self):
        self.assertEqual(motor.totales(1108765, 'FA'), (1108765, 0, 210665, 1319430))
        self.assertEqual(motor.totales(1000000, 'FE'), (0, 1000000, 0, 1000000))
        self.assertEqual(motor.totales(1000000, 'NC'), (0, 1000000, 0, 1000000))

    def test_glosa(self):
        self.assertEqual(motor.glosa(2026, 10, D('41065.38'), date(2026, 10, 1)), 'Octubre 2026 · Valor de la UF 41.065,38 del 01/10/2026')


def peso_uf(uf_, valor):
    return motor.peso(D(str(uf_)) * D(valor))


class FlujoTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.soc = Sociedad.objects.create(nombre='Inmobiliaria Prueba SpA', rut='76.000.000-0')
        cls.prop = Propiedad.objects.create(sociedad=cls.soc, codigo='prueba', nombre='Centro Prueba', nombre_glosa='Centro Prueba Talca')
        arr = Arrendatario.objects.create(razon_social='Tienda Ñandú SpA', rut='12.345.678-5')
        for cod, monto in (('P-01', '27'), ('P-02', '10')):
            c = Contrato.objects.create(codigo=cod, propiedad=cls.prop, sociedad=cls.soc, arrendatario=arr, nombre_fantasia=cod,
                                        inicio_vigencia=date(2025, 1, 1))
            TramoRenta.objects.create(contrato=c, desde=date(2025, 1, 1), monto=D(monto), unidad='UF')
        v = Contrato.objects.create(codigo='P-03', propiedad=cls.prop, sociedad=cls.soc, arrendatario=arr, nombre_fantasia='Variable',
                                    inicio_vigencia=date(2025, 1, 1), tipo_renta='MINVAR', pct_variable=D('0.02'), base_variable='V')
        TramoRenta.objects.create(contrato=v, desde=date(2025, 1, 1), monto=D('100'), unidad='UF')
        for f, val in UFS.items():
            UF.objects.create(fecha=f, valor=val)
        CodigoServicioBsoft.objects.create(sociedad=cls.soc, origen='RENTA', codigo='1002')
        CodigoServicioBsoft.objects.create(sociedad=cls.soc, origen='RDIF', codigo='1003')

    def test_borradores_y_unicidad(self):
        r = servicios.generar_borradores(self.soc, 2026, 10)
        self.assertEqual(len(r.creados), 3)
        self.assertTrue(any(cod == 'P-03' for cod, _ in r.con_advertencias))  # faltan ventas de septiembre
        servicios.generar_borradores(self.soc, 2026, 10)  # regenerar reemplaza, no duplica
        self.assertEqual(Documento.objects.filter(periodo=date(2026, 10, 1)).count(), 3)
        d = Documento.objects.get(contrato__codigo='P-01')
        self.assertEqual((d.neto, d.iva, d.total), (1108765, 210665, 1319430))

    def test_aprobacion_bloquea_advertencias(self):
        servicios.generar_borradores(self.soc, 2026, 10)
        ok, bloq = servicios.aprobar(self.soc, 2026, 10, 'test')
        self.assertEqual(sorted(ok), ['P-01', 'P-02'])
        self.assertEqual([c for c, _ in bloq], ['P-03'])
        VentaMensual.objects.create(contrato=Contrato.objects.get(codigo='P-03'), periodo=date(2026, 9, 1), monto_neto=900000000)
        servicios.generar_borradores(self.soc, 2026, 10, ['P-03'])
        ok, bloq = servicios.aprobar(self.soc, 2026, 10, 'test', ['P-03'])
        self.assertEqual(ok, ['P-03'])
        d = Documento.objects.get(contrato__codigo='P-03')
        self.assertEqual([l.origen for l in d.lineas.all()], ['RENTA', 'RDIF'])

    def test_aprobado_no_se_regenera(self):
        servicios.generar_borradores(self.soc, 2026, 10)
        servicios.aprobar(self.soc, 2026, 10, 'test', ['P-01'])
        r = servicios.generar_borradores(self.soc, 2026, 10, ['P-01'])
        self.assertEqual(r.creados, [])
        self.assertEqual(r.omitidos[0][0], 'P-01')

    def test_xml_bsoft(self):
        servicios.generar_borradores(self.soc, 2026, 10, ['P-01'])
        servicios.aprobar(self.soc, 2026, 10, 'test')
        with tempfile.TemporaryDirectory() as tmp:
            carpeta, z, archivos = bsoft.exportar_xml(self.soc, 2026, 10, carpeta=tmp)
            xml = open(archivos[0], newline='').read()
        self.assertIn('<RUTcliente>12345678-5</RUTcliente>', xml)
        self.assertIn('Centro Prueba Talca Valor de la UF 41.065,38 del 01/10/2026 Octubre 2026', xml)
        self.assertIn('<CodServicio>1002</CodServicio>', xml)
        self.assertIn('<Valor>1108765</Valor>', xml)
        self.assertIn('\r\n', xml)

    def test_xml_exige_codigos(self):
        CodigoServicioBsoft.objects.filter(origen='RENTA').delete()
        servicios.generar_borradores(self.soc, 2026, 10, ['P-01'])
        servicios.aprobar(self.soc, 2026, 10, 'test')
        with self.assertRaises(bsoft.FaltanCodigos):
            bsoft.exportar_xml(self.soc, 2026, 10)

    def test_rut_bsoft(self):
        self.assertEqual(bsoft.rut_bsoft('76.388.311-6'), '76388311-6')
        self.assertEqual(bsoft.rut_bsoft('12.345.678-k'), '12345678-K')
        with self.assertRaises(ValueError):
            bsoft.rut_bsoft('25.735.093-2 (cédula de extranjero)')


# ----------------------------------------------------------------------------- luz, gas y otros cargos
import os
import openpyxl

from facturacion import cargas, lectores
from facturacion.models import ConsumoServicio, MedidorServicio


def _planilla_luz(ruta):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Consumo Electrico '
    ws['C4'] = 'Septiembre'
    enc = ['Nº Local', 'nombre', 'mes', 'tarifa', 'cargo fijo ', 'Agosto', 'Septiembre', 'Cte. de factor', 'Consumos kw7h', 'x', 'y', 'z', 'a', 'b', 'A Facturar']
    for i, h in enumerate(enc):
        ws.cell(8, 3 + i, h)
    datos = [('Local 1B', 'Ale Cosmeticos', 'Septiembre', 'BT 1', 894.5, 12140, 12385, 1, 245, 0, 0, 0, 0, 0, 55360.5168),
             ('Local 5', 'Sin medidor', 'Septiembre', 'BT 1', 894.5, 100, 110, 1, 10, 0, 0, 0, 0, 0, 3000.4),
             ('Local 9', 'Vacío', 'Septiembre', 'BT 1', 894.5, 100, 100, 1, 0, 0, 0, 0, 0, 0, 0)]
    for r, fila in enumerate(datos, start=9):
        for i, v in enumerate(fila):
            ws.cell(r, 3 + i, v)
    ws.cell(12, 11, 255)  # fila de totales: sin nombre → termina la tabla
    wb.save(ruta)


class CargasTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.soc = Sociedad.objects.create(nombre='Inmobiliaria Prueba SpA', rut='76.000.000-0')
        prop = Propiedad.objects.create(sociedad=cls.soc, codigo='prueba', nombre='Centro Prueba')
        arr = Arrendatario.objects.create(razon_social='Ale SpA', rut='11.111.111-1')
        cls.c = Contrato.objects.create(codigo='P-01', propiedad=prop, sociedad=cls.soc, arrendatario=arr, nombre_fantasia='Ale',
                                        inicio_vigencia=date(2025, 1, 1))
        TramoRenta.objects.create(contrato=cls.c, desde=date(2025, 1, 1), monto=D('27'), unidad='UF')
        UF.objects.create(fecha=date(2026, 10, 1), valor=D('41065.38'))
        cls.med = MedidorServicio.objects.create(contrato=cls.c, servicio='ELEC', etiqueta_local='Local 1B', nombre_planilla='Ale Cosmeticos')

    def test_lector_electricidad(self):
        import tempfile
        with tempfile.TemporaryDirectory() as t:
            ruta = os.path.join(t, 'luz.xlsx')
            _planilla_luz(ruta)
            mes, filas = lectores.leer_electricidad(ruta)
        self.assertEqual(mes, 'Septiembre')
        self.assertEqual([f.nombre for f in filas], ['Ale Cosmeticos', 'Sin medidor', 'Vacío'])
        self.assertEqual(filas[0].cantidad, D('245'))

    def test_carga_asigna_ignora_ceros_y_reporta_sin_medidor(self):
        import tempfile
        with tempfile.TemporaryDirectory() as t:
            ruta = os.path.join(t, 'luz.xlsx')
            _planilla_luz(ruta)
            _, filas = lectores.leer_electricidad(ruta)
        res = cargas.cargar_filas(filas, date(2026, 9, 1), date(2026, 10, 1), 'luz.xlsx')
        self.assertEqual(res.cargados, 1)
        self.assertEqual(res.en_cero, ['Vacío'])
        self.assertEqual(res.sin_medidor[0][0], 'Sin medidor')
        self.assertEqual(ConsumoServicio.objects.get().neto, 55361)
        # recargar la misma planilla reemplaza, no duplica
        cargas.cargar_filas(filas, date(2026, 9, 1), date(2026, 10, 1), 'luz.xlsx')
        self.assertEqual(ConsumoServicio.objects.count(), 1)

    def test_factura_con_renta_luz_y_otro_y_aprobacion(self):
        ConsumoServicio.objects.create(contrato=self.c, servicio='ELEC', medidor='Local 1B', medidor_ref=self.med,
                                       periodo_consumo=date(2026, 9, 1), periodo_cobro=date(2026, 10, 1), lectura_anterior=12140,
                                       lectura_actual=12385, factor=1, cantidad=245, unidad_consumo='kWh', neto=55361)
        ConsumoServicio.objects.create(contrato=self.c, servicio='OTRO', concepto='reposicion llave', detalle='Llave bodega',
                                       periodo_consumo=date(2026, 10, 1), periodo_cobro=date(2026, 10, 1), neto=15000)
        servicios.generar_borradores(self.soc, 2026, 10)
        d = Documento.objects.get()
        self.assertEqual([l.origen for l in d.lineas.all()], ['RENTA', 'ELEC', 'OTRO'])
        self.assertEqual(d.neto, 1108765 + 55361 + 15000)
        self.assertEqual(d.iva, motor.peso(d.neto * D('0.19')))
        self.assertIn('lectura 12.140 → 12.385', d.lineas.get(origen='ELEC').detalle)
        self.assertEqual(d.lineas.get(origen='OTRO').concepto, 'REPOSICION LLAVE')
        # medidor sin confirmar → advertencia → no se aprueba solo
        self.assertTrue(any('por confirmar' in a for a in d.advertencias))
        ok, bloq = servicios.aprobar(self.soc, 2026, 10, 'test')
        self.assertEqual(ok, [])
        self.med.confirmado = True
        self.med.save()
        servicios.generar_borradores(self.soc, 2026, 10)
        ok, bloq = servicios.aprobar(self.soc, 2026, 10, 'test')
        self.assertEqual(ok, ['P-01'])

    def test_formato_numero_chileno(self):
        self.assertEqual(servicios._num(12140, 0), '12.140')
        self.assertEqual(servicios._num(945.35), '945,35')
        self.assertEqual(servicios._num(3.65, 3), '3,65')
