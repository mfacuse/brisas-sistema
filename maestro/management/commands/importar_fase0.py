"""Carga el maestro (sociedades, locales, arrendatarios, contratos, rentas, saldos) desde las planillas v2 de la Fase 0.

Uso:
    python manage.py importar_fase0                  # crea lo que falta; no toca contratos ya existentes
    python manage.py importar_fase0 --actualizar     # reemplaza datos y tramos de los contratos existentes
"""
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from maestro import historial, models as m

D = Decimal
SOCIEDADES = {
    'Boulevard': dict(nombre='Inmobiliaria Boulevard Talca SpA', rut='76.388.311-6'),
    'Plaza Motor': dict(nombre='Inmobiliaria 15 Oriente SpA', rut='76.602.726-1'),
}
PROPIEDADES = {
    'Boulevard': dict(codigo='boulevard', nombre='Brisas del Centro Boulevard', nombre_glosa='Brisas del Centro Boulevard Talca'),
    'Plaza Motor': dict(codigo='plaza-motor', nombre='Plaza Motor'),
}
TIPO_LOCAL = {'Local': 'LOCAL', 'Acceso peatonal': 'ACCESO', 'Módulo': 'MODULO', 'Supermercado': 'SUPER',
              'Publicidad': 'PUBLI', 'Estacionamiento': 'ESTAC'}

# Reglas de cobro que las planillas no traen en columnas propias (salen de los contratos y de las respuestas 1–33).
# Cualquier contrato no listado: renta fija, factura afecta, UF del día 1, prorrata por días.
CONFIG = {
    'B-01': dict(regla_uf='PH'),  # UF del primer día hábil (cláusula del contrato Mundo)
    'B-21': dict(tipo_renta='MINVAR', pct_variable=D('0.11'), base_variable='V'),
    'B-36': dict(tipo_renta='MINVAR', pct_variable=D('0.02'), base_variable='V'),
    'B-40': dict(tipo_renta='MINVAR', pct_variable=D('0.50'), base_variable='I', variable_desde=date(2027, 1, 1)),
    'B-39': dict(periodicidad='S', meses_cobro='1,6', reajuste_ipc=True),
}
FIN_ABIERTO = date(2099, 12, 31)


def f(v):
    """datetime → date; el resto igual."""
    return v.date() if isinstance(v, datetime) else v


def dec(v):
    return None if v in (None, '', '—') else D(str(v))


def txt(v):
    return '' if v in (None, '—') else str(v).strip()


def tabla(ws, fila_enc):
    """{encabezado: índice} de la fila de encabezados."""
    return {str(c.value).strip(): i for i, c in enumerate(ws[fila_enc]) if c.value is not None}


def col(enc, prefijo):
    for k, i in enc.items():
        if k.startswith(prefijo):
            return i
    raise CommandError(f'No encontré la columna «{prefijo}» (¿cambió la planilla?)')


class Command(BaseCommand):
    help = __doc__

    def add_arguments(self, p):
        base = Path(settings.BASE_DIR) / 'datos'
        p.add_argument('--contratos', default=str(base / 'Revision_Contratos_Brisas_2026-10-09_v2.xlsx'))
        p.add_argument('--maestro', default=str(base / 'Maestro_Locales_2026-10-09_v3.xlsx'))
        p.add_argument('--actualizar', action='store_true', help='reemplaza los contratos ya existentes')

    def handle(self, *a, **o):
        with historial.con_origen('importador'):
            self._importar(o)

    @transaction.atomic
    def _importar(self, o):
        for k in ('contratos', 'maestro'):
            if not Path(o[k]).exists():
                raise CommandError(f'No existe {o[k]}')
        self.actualizar = o['actualizar']
        wbc = openpyxl.load_workbook(o['contratos'], data_only=True)
        wbm = openpyxl.load_workbook(o['maestro'], data_only=True)
        self.n = dict(contratos=0, omitidos=0, tramos=0, locales=0, arrendatarios=0, saldos=0)
        self.sociedades = {}
        self.propiedades = {}
        for k, v in SOCIEDADES.items():
            s, _ = m.Sociedad.objects.update_or_create(rut=v['rut'], defaults={'nombre': v['nombre']})
            self.sociedades[k] = s
            pr = PROPIEDADES[k]
            self.propiedades[k], _ = m.Propiedad.objects.update_or_create(
                codigo=pr['codigo'], defaults={'sociedad': s, 'nombre': pr['nombre'], 'nombre_glosa': pr.get('nombre_glosa', '')})
        locales_de = self._locales(wbm)
        creados = self._contratos(wbc['Contratos'], locales_de)
        self._tramos(wbc['Renta_calendario'], creados)
        self._saldos(wbc['Saldos'], creados)
        self._servicios(wbc['Servicios'])
        self.stdout.write(self.style.SUCCESS(
            f"Listo: {self.n['contratos']} contratos, {self.n['tramos']} tramos de renta, {self.n['locales']} locales, "
            f"{self.n['arrendatarios']} arrendatarios nuevos, {self.n['saldos']} saldos. "
            f"Contratos que ya existían y se dejaron igual: {self.n['omitidos']}."))

    # ------------------------------------------------------------------ locales
    def _locales(self, wbm):
        """Crea los locales y devuelve {id_contrato: [Local]}."""
        out = {}
        for hoja, prop in (('Boulevard', 'Boulevard'), ('Plaza_Motor', 'Plaza Motor')):
            ws = wbm[hoja]
            enc = tabla(ws, 4)
            vistos = set()
            c_cod, c_tipo = 0, col(enc, 'Tipo')
            c_ct = col(enc, 'Contrato vigente')
            c_sit = col(enc, 'Situación')
            c_obs = col(enc, 'Observaciones')
            if hoja == 'Boulevard':
                c_piso, c_muni = col(enc, 'Piso'), col(enc, 'N° municipal')
                c_m2, c_calle, c_desig = col(enc, 'm² oficiales'), col(enc, 'Calle'), col(enc, 'Designación')
                c_cod = col(enc, 'Local (plano)')
            else:
                c_piso, c_muni = None, col(enc, 'N° en los contratos')
                c_m2, c_calle, c_desig = col(enc, 'm² as built'), col(enc, 'Dirección'), col(enc, 'Recinto')
            for r in ws.iter_rows(min_row=5, values_only=True):
                if not r[c_cod]:
                    continue
                cod = str(r[c_cod]).strip()
                ct = txt(r[c_ct])
                muni = txt(r[c_muni])
                if cod in vistos:  # p. ej. los tres «Acceso peatonal» o los dos módulos del pasillo central
                    cod = f'{cod} · {ct or muni}'
                vistos.add(cod)
                loc, nuevo = m.Local.objects.update_or_create(
                    propiedad=self.propiedades[prop], codigo_plano=cod,
                    defaults=dict(tipo=TIPO_LOCAL.get(txt(r[c_tipo]), 'LOCAL'), piso=txt(r[c_piso]) if c_piso is not None else '',
                                  calle=txt(r[c_calle]) if hoja == 'Boulevard' else '', numero_municipal=muni,
                                  designacion=txt(r[c_desig]), m2_oficial=dec(r[c_m2]),
                                  disponible=txt(r[c_sit]).startswith('Sin contrato (disponible'), observaciones=txt(r[c_obs])))
                self.n['locales'] += nuevo
                if ct:
                    out.setdefault(ct, []).append(loc)
        return out

    # ------------------------------------------------------------------ contratos
    def _contratos(self, ws, locales_de):
        enc = tabla(ws, 4)
        g = lambda r, p: r[col(enc, p)]
        creados = {}
        for r in ws.iter_rows(min_row=5, values_only=True):
            cod = r[0]
            if not (isinstance(cod, str) and re.fullmatch(r'[BO]-\d\d', cod)):
                continue
            prop_nombre = g(r, 'Propiedad')
            prop = self.propiedades[prop_nombre]
            existente = m.Contrato.objects.filter(codigo=cod).first()
            if existente and not self.actualizar:
                self.n['omitidos'] += 1
                continue
            rut_txt = txt(g(r, 'RUT'))
            mt = re.match(r'^([\d.]+-[\dkK])\s*(.*)$', rut_txt)
            rut_limpio, rut_nota = (mt.group(1), mt.group(2).strip('() ')) if mt else (rut_txt, '')
            arr, nuevo = m.Arrendatario.objects.get_or_create(rut=rut_limpio, razon_social=txt(g(r, 'Arrendatario')),
                                                              defaults={'representante': txt(g(r, 'Representante'))})
            self.n['arrendatarios'] += nuevo
            unidad_gar = g(r, 'Garantía (UF)'), g(r, 'Garantía (CLP)')
            datos = dict(
                propiedad=prop, sociedad=prop.sociedad, arrendatario=arr, nombre_fantasia=txt(g(r, 'Nombre de fantasía')),
                direccion_contrato=txt(g(r, 'Dirección según contrato')), representante=txt(g(r, 'Representante')),
                fecha_contrato=f(g(r, 'Fecha del contrato')), inicio_vigencia=f(g(r, 'Inicio de vigencia')),
                plazo_meses=int(g(r, 'Plazo inicial') or 12), fin_pactado=f(g(r, 'Fin pactado')),
                prorroga_meses=int(g(r, 'Prórroga') or 0), aviso_dias=int(g(r, 'Aviso')) if g(r, 'Aviso') not in (None, '') else None,
                m2_contrato=dec(g(r, 'm² (contrato)')), garantia_uf=dec(unidad_gar[0]), garantia_clp=dec(unidad_gar[1]),
                gastos_comunes=txt(g(r, 'Gastos comunes')), condiciones=txt(g(r, 'Condiciones especiales')),
                observaciones=(txt(g(r, 'Observaciones')) + (f' [Documento de identidad: {rut_nota}]' if rut_nota else '')).strip(), documentos_fuente=txt(g(r, 'Documentos fuente')),
                confianza_lectura=txt(g(r, 'Confianza de lectura')),
                tipo_renta='FIJA', periodicidad='M', meses_cobro='', reajuste_ipc=False, pct_variable=None, base_variable='',
                variable_desde=None, regla_uf='D1', modo_cambio_tramo='PR', tipo_documento='FA')
            datos.update(CONFIG.get(cod, {}))
            c, _ = m.Contrato.objects.update_or_create(codigo=cod, defaults=datos)
            c.tramos.all().delete()
            c.saldos.all().delete()
            m.ContratoLocal.objects.filter(contrato=c).delete()
            for loc in locales_de.get(cod, []):
                m.ContratoLocal.objects.create(contrato=c, local=loc)
            creados[cod] = c
            self.n['contratos'] += 1
        return creados

    # ------------------------------------------------------------------ tramos de renta
    def _tramos(self, ws, creados):
        enc = tabla(ws, 4)
        c_nota = col(enc, 'Nota')
        for r in ws.iter_rows(min_row=5, values_only=True):
            c = creados.get(r[0])
            if not c:
                continue
            desde, hasta = f(r[col(enc, 'Desde')]), f(r[col(enc, 'Hasta')])
            unidad = str(r[col(enc, 'Unidad')]).split('/')[0]
            nota = txt(r[c_nota])
            m.TramoRenta.objects.create(
                contrato=c, desde=desde, hasta=None if (hasta is None or hasta >= FIN_ABIERTO) else hasta,
                monto=dec(r[col(enc, 'Monto')]), unidad=unidad, tipo=txt(r[col(enc, 'Tipo de renta')]),
                fuente=txt(r[col(enc, 'Fuente')]), nota=nota, inferido='inferid' in nota.lower())
            self.n['tramos'] += 1

    # ------------------------------------------------------------------ saldos y servicios
    def _saldos(self, ws, creados):
        """Saldos solo de los contratos creados o actualizados en esta corrida (los demás no se tocan)."""
        enc = tabla(ws, 4)
        for r in ws.iter_rows(min_row=5, values_only=True):
            if not (isinstance(r[0], str) and re.fullmatch(r'[BO]-\d\d', r[0])):
                continue
            c = creados.get(r[0])
            if not c:
                continue
            c.saldos.all().delete()
            m.SaldoInformado.objects.create(
                contrato=c, concepto=txt(r[col(enc, 'Concepto')]) or 'Saldo adeudado', monto=dec(r[col(enc, 'Monto')]),
                unidad=txt(r[col(enc, 'Unidad')]), fecha=f(r[col(enc, 'Fecha informada')]),
                fuente=txt(r[col(enc, 'Fuente')]), nota=txt(r[col(enc, 'Nota')])[:200])
            self.n['saldos'] += 1

    def _servicios(self, ws):
        kv = {str(r[0]).strip(): (txt(r[1]), txt(r[2])) for r in ws.iter_rows(min_row=5, values_only=True) if r[0]}
        if 'ID' not in kv:
            return
        parte = kv.get('Parte contratada', ('', ''))[0]
        rut = re.search(r'RUT\s+([\d.]+-[\dkK])', parte)
        cond = '\n'.join(f'{k}: {v[0]}' for k, v in kv.items() if k not in ('ID', 'Contrato', 'Parte contratada', 'Parte contratante', 'Fecha del contrato', 'Vigencia'))
        m.ContratoServicio.objects.update_or_create(
            codigo=kv['ID'][0], defaults=dict(
                sociedad=self.sociedades['Boulevard'], contraparte='Megaservicios Talca SpA',
                rut_contraparte=rut.group(1) if rut else '', objeto=kv.get('Contrato', ('', ''))[0][:200],
                vigencia=kv.get('Vigencia', ('', ''))[0][:200], condiciones=cond))
